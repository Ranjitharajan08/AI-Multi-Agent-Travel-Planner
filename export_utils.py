from docx import Document

def export_travel_plan(plan_data, filename="travel_plan.docx"):
    doc = Document()

    doc.add_heading("Travel Itinerary", 0)

    # Basic Info
    doc.add_heading("Trip Summary", level=1)
    doc.add_paragraph(f"Destination: {plan_data['destination']}")
    doc.add_paragraph(f"Duration: {plan_data['duration']}")
    doc.add_paragraph(f"Budget: {plan_data['budget']}")

    # Itinerary
    for day in plan_data["itinerary"]:
        doc.add_heading(day["day"], level=2)

        for item in day["points"]:
            doc.add_paragraph(item, style='List Bullet')

    doc.save(filename)
    return filename